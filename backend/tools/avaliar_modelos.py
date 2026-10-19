"""Avalia modelos de sentimento contra posts rotulados à mão.

Diferente de ``benchmark_modelos.py`` (que mede o acerto sobre o *preço*, um
alvo muito ruidoso), aqui o alvo é o sentimento correto de cada texto. É a forma
rigorosa de decidir qual modelo usar.

Uso
---
    python tools/avaliar_modelos.py --csv rotulagem/posts_para_rotular.xlsx
    python tools/avaliar_modelos.py --csv rotulagem/amostra_sanidade.csv --base-real
    python tools/avaliar_modelos.py --csv a.csv --modelos ProsusAI/finbert

O CSV tem as colunas ``texto``, ``rotulo_a`` e (opcional) ``rotulo_b``. Com as
duas, o script mostra a concordância entre as pessoas (kappa de Cohen) e usa só
os posts em que elas concordaram. ``--base-real`` mostra também como cada modelo
distribui os posts reais do banco (um modelo enviesado marca quase tudo igual).
"""

import argparse
import csv
import math
import sys
import time
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from services import sentimento as svc_sentimento

CLASSES = ("negativo", "neutro", "positivo")

MODELOS_PADRAO = [
    "nlptown/bert-base-multilingual-uncased-sentiment",
    "ProsusAI/finbert",
    "cardiffnlp/twitter-roberta-base-sentiment-latest",
    "ElKulako/cryptobert",
]

_ALIAS = {
    "positivo": "positivo", "pos": "positivo", "p": "positivo", "+": "positivo",
    "negativo": "negativo", "neg": "negativo", "n": "negativo", "-": "negativo",
    "neutro": "neutro", "neu": "neutro", "0": "neutro",
}


def normalizar_rotulo(valor: str | None) -> str | None:
    return _ALIAS.get((valor or "").strip().lower())


def ler_csv(caminho: Path) -> list[dict]:
    """Lê a planilha de rótulos (.xlsx ou .csv com separador ';')."""
    if caminho.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(caminho, read_only=True, data_only=True)["Posts"]
        linhas = ws.iter_rows(values_only=True)
        cabecalho = [str(c or "").strip() for c in next(linhas)]
        return [
            {k: ("" if v is None else str(v)) for k, v in zip(cabecalho, ln)}
            for ln in linhas
        ]
    with open(caminho, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f, delimiter=";"))


def kappa_cohen(a: list[str], b: list[str]) -> float | None:
    """Concordância entre dois rotuladores, corrigida pelo acaso (-1 a 1)."""
    n = len(a)
    if n == 0:
        return None
    observada = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    esperada = sum((ca[c] / n) * (cb[c] / n) for c in CLASSES)
    if esperada == 1:
        return None
    return (observada - esperada) / (1 - esperada)


def intervalo_wilson(acertos: int, n: int, z: float = 1.96) -> tuple[float, float]:
    if n == 0:
        return (0.0, 0.0)
    p = acertos / n
    denom = 1 + z * z / n
    centro = (p + z * z / (2 * n)) / denom
    margem = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denom
    return (max(0.0, centro - margem) * 100, min(1.0, centro + margem) * 100)


def metricas(verdadeiros: list[str], previstos: list[str]) -> dict:
    n = len(verdadeiros)
    acertos = sum(v == p for v, p in zip(verdadeiros, previstos))
    f1s = {}
    for c in CLASSES:
        tp = sum(v == c and p == c for v, p in zip(verdadeiros, previstos))
        fp = sum(v != c and p == c for v, p in zip(verdadeiros, previstos))
        fn = sum(v == c and p != c for v, p in zip(verdadeiros, previstos))
        prec = tp / (tp + fp) if tp + fp else 0.0
        rec = tp / (tp + fn) if tp + fn else 0.0
        f1s[c] = 2 * prec * rec / (prec + rec) if prec + rec else 0.0
    confusao = {(v, p): 0 for v in CLASSES for p in CLASSES}
    for v, p in zip(verdadeiros, previstos):
        confusao[(v, p)] += 1
    return {
        "n": n,
        "acuracia": acertos / n if n else 0.0,
        "ic": intervalo_wilson(acertos, n),
        "f1_macro": sum(f1s.values()) / len(CLASSES),
        "f1": f1s,
        "confusao": confusao,
    }


def classificar(nome_modelo: str, textos: list[str]) -> tuple[list[str], float]:
    from transformers import pipeline

    pipe = pipeline("sentiment-analysis", model=nome_modelo, tokenizer=nome_modelo)
    id2label = getattr(pipe.model.config, "id2label", None)
    inicio = time.time()
    saidas = pipe([t[:2000] for t in textos], truncation=True, max_length=512,
                  batch_size=8)
    return ([svc_sentimento.mapear_label(s["label"], id2label) for s in saidas],
            time.time() - inicio)


def posts_reais() -> list[str]:
    from db import SessionLocal
    from models import SocialPost

    db = SessionLocal()
    try:
        textos = [p.texto or "" for p in db.query(SocialPost).all()]
    finally:
        db.close()
    return [t for t in textos if svc_sentimento.texto_e_crypto_relevante(t)]


def carregar_rotulados(caminhos: list[str]) -> tuple[list[str], list[str]]:
    textos, rotulos = [], []
    for c in caminhos:
        linhas = ler_csv(Path(c))
        pares = []
        for ln in linhas:
            a = normalizar_rotulo(ln.get("rotulo_a"))
            b = normalizar_rotulo(ln.get("rotulo_b"))
            if a and "rotulo_b" in ln and (ln.get("rotulo_b") or "").strip():
                pares.append((ln["texto"], a, b))
            elif a:
                pares.append((ln["texto"], a, None))
        dois = [(a, b) for _, a, b in pares if b]
        if dois:
            k = kappa_cohen([a for a, _ in dois], [b for _, b in dois])
            conc = sum(a == b for a, b in dois) / len(dois)
            print(f"{c}: {len(dois)} posts com dois rotuladores, concordância "
                  f"{conc:.0%}, kappa {k:.2f}" if k is not None else
                  f"{c}: {len(dois)} posts com dois rotuladores")
        for texto, a, b in pares:
            if b is None or a == b:
                textos.append(texto)
                rotulos.append(a)
    return textos, rotulos


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--csv", nargs="+", required=True)
    parser.add_argument("--modelos", nargs="+", default=MODELOS_PADRAO)
    parser.add_argument("--base-real", action="store_true",
                        help="mostra a distribuição de cada modelo no banco")
    args = parser.parse_args()

    textos, rotulos = carregar_rotulados(args.csv)
    if not textos:
        sys.exit("Nenhum post rotulado encontrado nos CSVs informados.")
    dist = Counter(rotulos)
    print(f"\n{len(textos)} posts de referência: "
          + ", ".join(f"{c} {dist[c]}" for c in CLASSES))
    reais = posts_reais() if args.base_real else []

    linhas = []
    for nome in args.modelos:
        print(f"\n>>> {nome}")
        try:
            previstos, seg = classificar(nome, textos)
            real_prev = classificar(nome, reais)[0] if reais else []
        except Exception as e:  # modelo indisponível, sem rede, etc.
            print(f"    falhou: {e}")
            continue
        m = metricas(rotulos, previstos)
        linhas.append((nome, m, seg, Counter(real_prev)))
        print(f"    acurácia {m['acuracia']:.1%}  F1 macro {m['f1_macro']:.3f}  "
              f"({seg:.1f}s)")
        print("    verdadeiro \\ previsto:  " + "  ".join(f"{c[:3]:>4}" for c in CLASSES))
        for v in CLASSES:
            print(f"      {v:<9}            "
                  + "  ".join(f"{m['confusao'][(v, p)]:>4}" for p in CLASSES))

    print("\n" + "=" * 78)
    print(f"{'Modelo':<50} {'Acur.':>6} {'IC 95%':>14} {'F1':>6}")
    for nome, m, _, _ in sorted(linhas, key=lambda x: -x[1]["f1_macro"]):
        lo, hi = m["ic"]
        print(f"{nome:<50} {m['acuracia']:>6.1%} {f'[{lo:.0f}% - {hi:.0f}%]':>14} "
              f"{m['f1_macro']:>6.3f}")
    if reais:
        print(f"\nDistribuição nos {len(reais)} posts reais do banco "
              "(modelo enviesado = quase tudo na mesma coluna):")
        print(f"{'Modelo':<50} {'neg':>6} {'neu':>6} {'pos':>6}")
        for nome, _, _, c in linhas:
            n = sum(c.values()) or 1
            print(f"{nome:<50} {c['negativo'] / n:>6.0%} {c['neutro'] / n:>6.0%} "
                  f"{c['positivo'] / n:>6.0%}")


if __name__ == "__main__":
    main()
